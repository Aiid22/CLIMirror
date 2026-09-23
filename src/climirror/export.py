"""Four-column, text-safe XLSX export."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AcceptedMapping, ea


HEADERS = ("Command", "Handler Address and Name", "Handler Relative Path", "Evidence Chain")


def _excel_text(value: str) -> str:
    if value and value[0] in "=+-@\t\r\n":
        return "'" + value
    return value


def write_workbook(path: Path, mappings: list[AcceptedMapping]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CLIMirror"
    sheet.append(HEADERS)
    seen: set[tuple[str, str, int]] = set()
    count = 0
    for item in sorted(mappings, key=lambda mapping: mapping.identity()):
        identity = item.identity()
        if identity in seen:
            continue
        seen.add(identity)
        handler_name = f"[IDA auto] {item.handler.name}" if item.handler.auto_name else item.handler.name
        values = (item.command, f"{ea(item.handler.address)} | {handler_name}", item.handler.relative_path, item.evidence_chain)
        sheet.append([_excel_text(str(value)) for value in values])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        count += 1
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16324F")
    for column, width in enumerate((46, 40, 45, 110), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    workbook.save(path)
    return count
